"""Approximate PNG preview of a pptx (geometry + fills + images + text).
Not a pixel-perfect renderer — it draws real shape boxes/fills/images and wraps
text using Arial metrics so layout, overlaps, and balance can be reviewed.
"""
import sys, io
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

PX = 96.0 / 914400.0  # EMU -> px at 96 dpi
def px(emu): return int(round(emu * PX))

FONT_DIR = r"C:\Windows\Fonts"
def font(size_pt, bold=False, italic=False):
    px_size = int(round(size_pt * 96.0 / 72.0))
    name = "arialbd.ttf" if bold else "arial.ttf"
    if italic and not bold: name = "ariali.ttf"
    try:
        return ImageFont.truetype(f"{FONT_DIR}\\{name}", px_size)
    except Exception:
        return ImageFont.truetype(f"{FONT_DIR}\\arial.ttf", px_size)

def rgb(color):
    try:
        return (color[0], color[1], color[2])
    except Exception:
        return None

def shape_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:  # solid
            c = sh.fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None

def shape_line(sh):
    try:
        if sh.line.color and sh.line.color.type is not None:
            c = sh.line.color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None

def wrap(draw, text, fnt, maxw):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines

def draw_text_frame(draw, sh, L, T, W, H):
    tf = sh.text_frame
    anchor = tf.vertical_anchor
    # build lines
    blocks = []
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            blocks.append((" ", 12, (0,0,0), PP_ALIGN.LEFT, False, False)); continue
        txt = "".join(r.text for r in runs)
        r0 = runs[0]
        size = r0.font.size.pt if r0.font.size else 14
        col = rgb(r0.font.color.rgb) if (r0.font.color and r0.font.color.type is not None) else (30,30,30)
        if col is None: col = (30,30,30)
        bold = bool(r0.font.bold); ital = bool(r0.font.italic)
        align = p.alignment or PP_ALIGN.LEFT
        blocks.append((txt, size, col, align, bold, ital))
    # measure
    rendered = []
    total_h = 0
    pad = 6
    for txt, size, col, align, bold, ital in blocks:
        fnt = font(size, bold, ital)
        lines = wrap(draw, txt, fnt, W - 2*pad)
        lh = int(size * 96/72 * 1.18)
        for ln in lines:
            rendered.append((ln, fnt, col, align, lh))
            total_h += lh
    if anchor == MSO_ANCHOR.MIDDLE:
        y = T + max(0, (H - total_h)//2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = T + max(0, H - total_h)
    else:
        y = T + pad
    for ln, fnt, col, align, lh in rendered:
        tw = draw.textlength(ln, font=fnt)
        if align == PP_ALIGN.CENTER: x = L + (W - tw)//2
        elif align == PP_ALIGN.RIGHT: x = L + W - tw - pad
        else: x = L + pad
        draw.text((x, y), ln, font=fnt, fill=col)
        y += lh

import re as _re
def render_layout(img, d, slide, W, H):
    """Approximate the inherited layout: background image + non-placeholder pictures (banner/logo)."""
    try:
        layout = slide.slide_layout
    except Exception:
        return
    # background blip
    try:
        xml = layout.part.blob.decode("utf8", "ignore")
        m = _re.search(r"<p:bg>.*?<a:blip[^>]*r:embed=\"(rId\d+)\"", xml, _re.S)
        if m:
            part = layout.part.related_part(m.group(1))
            bg = Image.open(io.BytesIO(part.blob)).convert("RGB").resize((W, H))
            img.paste(bg, (0, 0))
    except Exception:
        pass
    # layout pictures (banner, logo) — skip placeholders
    for sh in layout.shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.left is not None:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
                im = im.resize((max(1, px(sh.width)), max(1, px(sh.height))))
                img.paste(im, (px(sh.left), px(sh.top)), im)
        except Exception:
            pass

def render(path, outdir):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    import os; os.makedirs(outdir, exist_ok=True)
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (255,255,255))
        d = ImageDraw.Draw(img)
        render_layout(img, d, slide, W, H)
        d.rectangle([0,0,W-1,H-1], outline=(220,220,220))
        for sh in slide.shapes:
            if sh.left is None: continue
            L,T,Wd,Hd = px(sh.left),px(sh.top),px(sh.width),px(sh.height)
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                    im = im.resize((max(1,Wd),max(1,Hd)))
                    img.paste(im, (L,T))
                    d.rectangle([L,T,L+Wd,T+Hd], outline=(150,160,150))
                except Exception as e:
                    d.rectangle([L,T,L+Wd,T+Hd], outline=(255,0,0))
                continue
            if sh.has_table:
                tbl = sh.table
                rows=len(tbl.rows); cols=len(tbl.columns)
                colw=[px(c.width) for c in tbl.columns]
                rowh=Hd//rows
                yy=T
                for ri in range(rows):
                    xx=L
                    for ci in range(cols):
                        cell=tbl.cell(ri,ci); cw=colw[ci]
                        f=None
                        try:
                            if cell.fill.type==1: cc=cell.fill.fore_color.rgb; f=(cc[0],cc[1],cc[2])
                        except Exception: pass
                        if f: d.rectangle([xx,yy,xx+cw,yy+rowh], fill=f)
                        d.rectangle([xx,yy,xx+cw,yy+rowh], outline=(180,180,180))
                        ct=cell.text
                        if ct.strip():
                            fnt=font(11)
                            for ln in wrap(d,ct,fnt,cw-6)[:2]:
                                d.text((xx+4,yy+4),ln,font=fnt,fill=(40,40,40))
                                yy2=0
                        xx+=cw
                    yy+=rowh
                continue
            fill = shape_fill(sh)
            line = shape_line(sh)
            if fill or line:
                d.rectangle([L,T,L+Wd,T+Hd], fill=fill, outline=line or fill)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(d, sh, L,T,Wd,Hd)
        op = os.path.join(outdir, f"slide_{i:02d}.png")
        img.save(op); paths.append(op)
    # contact sheet
    cols=3; rows=(len(paths)+cols-1)//cols
    tw,th=W//3, H//3
    sheet=Image.new("RGB",(tw*cols, th*rows),(245,245,245))
    for idx,pp in enumerate(paths):
        im=Image.open(pp).resize((tw,th))
        sheet.paste(im,((idx%cols)*tw,(idx//cols)*th))
    sheet.save(os.path.join(outdir,"_contact_sheet.png"))
    print("rendered", len(paths), "slides ->", outdir)

if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
